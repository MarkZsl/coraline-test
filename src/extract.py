from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

from src.logger import get_logger

logger = get_logger(__name__)

EXPECTED_HEADER = ("ID", "Date", "Region", "City", "Category", "Product", "Qty", "UnitPrice", "TotalPrice")


def _is_blank_row(row: tuple) -> bool:
    return all(cell is None for cell in row)


def _is_header_row(row: tuple) -> bool:
    return tuple(row[: len(EXPECTED_HEADER)]) == EXPECTED_HEADER


def extract_food_sales(file_path: str | Path, sheet_name: str = "FoodSales") -> pd.DataFrame:
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"Excel source file not found: {file_path}")

    logger.info("Reading workbook '%s' (sheet='%s')", file_path, sheet_name)
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")
    ws = wb[sheet_name]

    blocks: list[list[tuple]] = []
    current_block: list[tuple] | None = None
    n_headers_found = 0

    for row in ws.iter_rows(values_only=True):
        if _is_header_row(row):
            current_block = []
            blocks.append(current_block)
            n_headers_found += 1
            continue

        if current_block is None:
            continue

        if _is_blank_row(row):
            current_block = None
            continue

        current_block.append(row[: len(EXPECTED_HEADER)])

    wb.close()

    if n_headers_found == 0:
        raise ValueError(
            f"No header row matching {EXPECTED_HEADER} was found in sheet '{sheet_name}'. "
            "The source file layout may have changed."
        )

    total_rows = sum(len(b) for b in blocks)
    logger.info("Found %d data block(s) in '%s' totaling %d rows", n_headers_found, sheet_name, total_rows)

    all_rows = [r for block in blocks for r in block]
    df = pd.DataFrame(all_rows, columns=EXPECTED_HEADER)

    df = _clean(df)
    logger.info("Extracted %d clean rows, columns=%s", len(df), list(df.columns))
    return df


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()

    df["ID"] = df["ID"].astype(str).str.strip()
    df["Date"] = pd.to_datetime(df["Date"]).dt.date
    for col in ("Region", "City", "Category", "Product"):
        df[col] = df[col].astype(str).str.strip()

    df["Qty"] = pd.to_numeric(df["Qty"], errors="coerce").astype("Int64")
    df["UnitPrice"] = pd.to_numeric(df["UnitPrice"], errors="coerce").round(2)
    df["TotalPrice"] = pd.to_numeric(df["TotalPrice"], errors="coerce").round(2)

    before = len(df)
    df = df.dropna(subset=["ID", "Date", "Qty", "UnitPrice", "TotalPrice"])
    dropped = before - len(df)
    if dropped:
        logger.warning("Dropped %d row(s) with missing required values", dropped)

    df = df.drop_duplicates(subset=["ID"], keep="first")

    # Rename to snake_case for Postgres convention.
    df = df.rename(
        columns={
            "ID": "id",
            "Date": "sale_date",
            "Region": "region",
            "City": "city",
            "Category": "category",
            "Product": "product",
            "Qty": "qty",
            "UnitPrice": "unit_price",
            "TotalPrice": "total_price",
        }
    )
    return df[["id", "sale_date", "region", "city", "category", "product", "qty", "unit_price", "total_price"]]
